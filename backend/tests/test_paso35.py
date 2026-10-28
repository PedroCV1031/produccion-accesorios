import io
import unittest
from openpyxl import Workbook
from app.services.lector_excel_acc import detectar_hojas_validas, analizar_archivo
from app.services.identificacion import clave_unica, conciliar_lote

class Paso35(unittest.TestCase):
    def libro(self):
        wb=Workbook(); w=wb.active;w.title='OCT 2026'
        encabezados=['FECHA EN LA QUE EL CLIENTE COLOCA EL PEDIDO','FECHA PARA LA QUE EL CLIENTE NECESITA EL DESPACHO','FECHA EN LA QUE ARCOLINE SE COMPROMETE A ENTREGAR','CLIENTE','COMERCIAL','ORDEN DE COMPRA','REFERENCIA','ACABADO / TALLA','COLOR Y CODIGO','CANTIDAD TOTAL']
        for i,h in enumerate(encabezados,1): w.cell(1,i,h)
        for i,h in enumerate(['FECHA ULTIMA ENTREGA','CANTIDAD ENTREGADA','FACTURA'],11):w.cell(2,i,h)
        w.append(['2026-10-01',None,None,'Cliente A','C',None,'99131','ZAMAC','NEGRO',5000,None,0,None])
        wb.create_sheet('Hoja1')
        out=io.BytesIO();wb.save(out);return out.getvalue()
    def test_seleccion_y_oc_vacia(self):
        data=self.libro(); r=detectar_hojas_validas(data)
        self.assertEqual([x['formato_valido'] for x in r],[True,False])
        a=analizar_archivo(data,hojas=['OCT 2026'],incluir_registros=True)
        self.assertEqual(a['total_validas'],1)
        self.assertEqual(a['total_invalidas'],0)
        self.assertEqual(a['registros'][0]['orden_compra'],None)
    def test_clave_vacia_y_duplicados(self):
        a={'cliente':'Cliente A','orden_compra':None,'referencia':'99131','acabado_talla':'ZAMAC','color_codigo':'NEGRO','fecha_pedido':'2026-10-01'}
        b={**a,'orden_compra':'   ','cantidad_solicitada':9000}
        self.assertEqual(clave_unica(a),clave_unica(b))
        self.assertEqual([x.accion for x in conciliar_lote([a,b],[])],['duplicado','duplicado'])
if __name__=='__main__':unittest.main()
