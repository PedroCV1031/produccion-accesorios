"""Lectura y validación preliminar de RADICADO ACC (sin escribir en MongoDB)."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

COLUMNAS = (
    "fecha_pedido", "fecha_solicitada", "fecha_comprometida", "cliente",
    "comercial", "orden_compra", "referencia", "acabado_talla", "color_codigo",
    "cantidad_solicitada", "fecha_ultima_entrega", "cantidad_entregada_acumulada", "factura",
)
MAX_FILAS = 200_000


def normalizar_encabezado(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().casefold()


def validar_encabezados(ws):
    esperados = ["FECHA EN LA QUE EL CLIENTE COLOCA EL PEDIDO",
                "FECHA PARA LA QUE EL CLIENTE NECESITA EL DESPACHO",
                "FECHA EN LA QUE ARCOLINE SE COMPROMETE A ENTREGAR",
                "CLIENTE", "COMERCIAL", "ORDEN DE COMPRA", "REFERENCIA",
                "ACABADO / TALLA", "COLOR Y CODIGO", "CANTIDAD TOTAL"]
    valores = [ws.cell(1, i).value for i in range(1, 11)]
    if [normalizar_encabezado(x) for x in valores] != [normalizar_encabezado(x) for x in esperados]:
        return False, "Encabezados A:J no coinciden con el formato ACC"
    detalle = [ws.cell(2, i).value for i in (11, 12, 13)]
    aliases = [
        {"fecha ultima entrega"},
        {"cantidad entregada", "cantidad entregada acumulada"},
        {"factura", "factura ultima entrega"},
    ]
    if any(normalizar_encabezado(x) not in accepted for x, accepted in zip(detalle, aliases)):
        return False, "Encabezados K2:M2 no coinciden con el bloque de entregas ACC"
    return True, None


def texto(valor):
    if valor is None: return None
    val = str(valor).strip()
    return val if val else None


def cantidad(valor):
    if valor is None or str(valor).strip() == "": return None
    if isinstance(valor, bool): raise ValueError("cantidad no numérica")
    try:
        if isinstance(valor, str):
            s = valor.strip().replace(" ", "")
            if "," in s and "." in s: s = s.replace(".", "").replace(",", ".")
            elif "," in s: s = s.replace(",", ".")
            elif s.count(".") == 1 and len(s.split(".")[1]) == 3: s = s.replace(".", "")
            n = Decimal(s)
        else: n = Decimal(str(valor))
    except (InvalidOperation, ValueError): raise ValueError("cantidad no numérica")
    if not n.is_finite() or n != n.to_integral_value(): raise ValueError("cantidad debe ser un entero")
    return int(n)


def fecha(valor, epoch):
    if valor is None or str(valor).strip() == "": return None
    if isinstance(valor, datetime): return valor.date().isoformat()
    if isinstance(valor, date): return valor.isoformat()
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try: return from_excel(valor, epoch=epoch).date().isoformat()
        except (ValueError, OverflowError, AttributeError): raise ValueError("fecha Excel inválida")
    if isinstance(valor, str):
        raw = valor.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try: return datetime.strptime(raw, fmt).date().isoformat()
            except ValueError: pass
    raise ValueError("fecha inválida")


def analizar_archivo(contenido: bytes, hojas: list[str] | None = None, limite_ejemplos: int = 15, *, incluir_registros: bool = False):
    if len(contenido) > 20 * 1024 * 1024: raise ValueError("Archivo mayor de 20 MB")
    try: wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True, keep_links=False)
    except Exception as exc: raise ValueError("No fue posible abrir el libro Excel") from exc
    try:
        seleccion = list(dict.fromkeys(hojas if hojas is not None else wb.sheetnames))
        if any(h not in wb.sheetnames for h in seleccion): raise ValueError("La selección contiene hojas inexistentes")
        resultados=[]; registros=[]
        for nombre in seleccion:
            ws=wb[nombre]
            correcto, razon=validar_encabezados(ws)
            if not correcto:
                resultados.append({"hoja": nombre, "formato_valido": False, "motivo": razon, "filas_validas":0,"filas_invalidas":0,"ejemplos":[]})
                continue
            buenas=malas=0; ejemplos=[]; errores=[]
            for numero_fila, valores in enumerate(ws.iter_rows(min_row=3, max_col=13, values_only=True), start=3):
                if buenas + malas > MAX_FILAS: raise ValueError("Límite de filas superado")
                if all(x is None or str(x).strip()=="" for x in valores): continue
                registro=dict(zip(COLUMNAS,valores))
                try:
                    for k in ("fecha_pedido","fecha_solicitada","fecha_comprometida","fecha_ultima_entrega"):
                        registro[k]=fecha(registro[k],wb.epoch)
                    for k in ("cliente","comercial","orden_compra","referencia","acabado_talla","color_codigo","factura"):
                        registro[k]=texto(registro[k])
                    for k in ("cantidad_solicitada","cantidad_entregada_acumulada"):
                        registro[k]=cantidad(registro[k])
                    from app.services.identificacion import CAMPOS_IDENTIDAD
                    if any(not registro.get(k) for k in CAMPOS_IDENTIDAD if k != "orden_compra"):
                        raise ValueError("Falta campo obligatorio de la clave: cliente, referencia, acabado/talla, color/código o fecha pedido")
                    if registro["cantidad_solicitada"] is None or registro["cantidad_solicitada"] <= 0: raise ValueError("cantidad solicitada debe ser positiva")
                    if registro["cantidad_entregada_acumulada"] is None: registro["cantidad_entregada_acumulada"]=0
                    if registro["cantidad_entregada_acumulada"] < 0: raise ValueError("cantidad entregada negativa")
                    registro["cantidad_pendiente"] = registro["cantidad_solicitada"]-registro["cantidad_entregada_acumulada"]
                    buenas+=1
                    if incluir_registros:
                        registros.append({**registro, "_hoja": nombre, "_fila": numero_fila})
                    if len(ejemplos)<limite_ejemplos: ejemplos.append(registro)
                except ValueError as exc:
                    malas+=1
                    errores.append({"fila": numero_fila, "motivo": str(exc), "datos_originales": {k: (v.isoformat() if isinstance(v, (datetime, date)) else v if isinstance(v, (str, int, float, bool, type(None))) else str(v)) for k, v in zip(COLUMNAS, valores)}})
            resultados.append({"hoja":nombre,"formato_valido":True,"filas_validas":buenas,"filas_invalidas":malas,"ejemplos":ejemplos,"errores":errores})
        salida = {"hojas_disponibles":wb.sheetnames,"hojas":resultados,"total_validas":sum(x["filas_validas"] for x in resultados),"total_invalidas":sum(x["filas_invalidas"] for x in resultados),"solo_vista_previa":True}
        if incluir_registros: salida["registros"] = registros
        return salida
    finally: wb.close()


def detectar_hojas_validas(contenido: bytes) -> list[dict]:
    """Solo revisa estructura, sin leer los miles de pedidos ni escribir en MongoDB."""
    if len(contenido) > 20 * 1024 * 1024:
        raise ValueError("Archivo mayor de 20 MB")
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("No fue posible abrir el libro Excel") from exc
    try:
        resultado = []
        for nombre in libro.sheetnames:
            valido, motivo = validar_encabezados(libro[nombre])
            resultado.append({"hoja": nombre, "formato_valido": valido, "motivo": motivo})
        return resultado
    finally:
        libro.close()
